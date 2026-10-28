import copy
import json
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from resume_system.config import load_style
from resume_system.export import docx_export as de

ROOT = Path(__file__).resolve().parents[1]


def _demo():
    kb = json.loads((ROOT / "examples/demo/data/knowledge_base/knowledge_base.json").read_text(encoding="utf-8"))
    task = json.loads(next((ROOT / "examples/demo/data/jobs").glob("*.json")).read_text(encoding="utf-8"))
    return task["versions"][-1]["resume_json"], kb


def _render(resume, kb, tmp_path):
    out = tmp_path / "x.docx"
    de._render_resume(resume, kb["profile"], out, {e["id"]: e for e in kb["experiences"]})
    return Document(str(out))


def _ki(kb):
    return {e["id"]: e for e in kb["experiences"]}


def test_default_theme_is_finance():
    s = load_style()
    assert s["colors"]["accent"] == "#1F3864"


def test_heading_uses_accent_color(tmp_path):
    resume, kb = _demo()
    doc = _render(resume, kb, tmp_path)
    accent = RGBColor.from_string(load_style()["colors"]["accent"].lstrip("#"))
    for p in doc.paragraphs:
        if p.text.strip() == "教育背景":
            assert p.runs[0].font.color.rgb == accent
            return
    raise AssertionError("没找到「教育背景」标题")


def test_layout_order_reorder(tmp_path):
    resume, kb = _demo()
    resume = copy.deepcopy(resume)
    resume["skills"] = ["Excel", "SQL", "Stata"]
    old_order, old_show = de.LAYOUT_ORDER, de.LAYOUT_SHOW
    try:
        de.LAYOUT_ORDER = ["abilities", "sections", "skills", "education"]
        de.LAYOUT_SHOW = {"photo": True, "availability": True, "job_title": True}
        doc = _render(resume, kb, tmp_path)
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "个人能力" in texts and "技能" in texts
        assert texts.index("个人能力") < texts.index("技能")
    finally:
        de.LAYOUT_ORDER, de.LAYOUT_SHOW = old_order, old_show


def test_show_toggles_hide(tmp_path):
    resume, kb = _demo()
    old_order, old_show = de.LAYOUT_ORDER, de.LAYOUT_SHOW
    try:
        de.LAYOUT_ORDER = ["education", "sections", "abilities"]
        de.LAYOUT_SHOW = {"photo": True, "availability": False, "job_title": False}
        doc = _render(resume, kb, tmp_path)
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "到岗时间" not in full
        assert "求职意向" not in full
    finally:
        de.LAYOUT_ORDER, de.LAYOUT_SHOW = old_order, old_show


def test_skills_render(tmp_path):
    resume, kb = _demo()
    resume = copy.deepcopy(resume)
    resume["skills"] = ["Excel", "SQL", "Stata", "Wind"]
    doc = _render(resume, kb, tmp_path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "技能" in texts
    assert any("Excel" in t and "Stata" in t for t in texts)


def test_skills_affect_height():
    resume, kb = _demo()
    resume = copy.deepcopy(resume)
    base = de._estimated_height(resume, kb["profile"], _ki(kb))
    resume["skills"] = ["Excel", "SQL", "Stata"]
    with_skills = de._estimated_height(resume, kb["profile"], _ki(kb))
    assert with_skills > base


def test_fit_to_one_page_deletes_bullets():
    resume, kb = _demo()
    r = copy.deepcopy(resume)
    bullet = {
        "label": "边界测试",
        "desc": "使用 Excel 与 SQL 维护订单台账并逐项核对轨迹一致性，独立完成数据比对与差异标注，输出结构化报表，支撑业务复盘与排期决策，体现数据分析与问题定位能力。",
    }
    r["sections"][0]["experiences"][0]["bullets"] += [copy.deepcopy(bullet) for _ in range(30)]
    ki = _ki(kb)
    fitted = de._fit_to_one_page(r, {}, kb["profile"], ki)
    assert de._fits_one_page(fitted, kb["profile"], ki)
    n_before = sum(len(e["bullets"]) for s in r["sections"] for e in s["experiences"])
    n_after = sum(len(e["bullets"]) for s in fitted["sections"] for e in s["experiences"])
    assert n_after < n_before
